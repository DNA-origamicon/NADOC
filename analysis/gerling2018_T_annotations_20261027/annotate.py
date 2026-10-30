"""Review-only unpaired-T hypotheses; never runs REVNANO or changes a design."""
from pathlib import Path
from functools import lru_cache
from collections import defaultdict,Counter
import json,hashlib,csv,html
import openpyxl
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.cell.rich_text import CellRichText,TextBlock
from openpyxl.cell.text import InlineFont
R=Path(__file__).resolve().parents[2];O=Path(__file__).resolve().parent
S=(R/'backend/core/p7560.txt').read_text().strip().upper();N=len(S)
C=S.translate(str.maketrans('ACGT','TGCA'))
seed=defaultdict(list)
for p in range(N):seed[''.join(C[(p-k)%N] for k in range(4))].append(p)
TOP=8

def annotate(q,key):
 n=len(q);runs=[]
 for i in range(n):
  rr=[]
  for p in seed.get(q[i:i+4],[]):
   k=4
   while i+k<n and q[i+k]==C[(p-k)%N]:k+=1
   rr.append((p,k))
  runs.append(rr)
 protected=set()
 for i,rr in enumerate(runs):
  longhits=[(p,k) for p,k in rr if k>=12]
  if len(longhits)==1:
   protected.update(range(i,i+longhits[0][1]))
 longest={}
 for i in range(n-1,-1,-1):
  maximum=max((k for p,k in runs[i]),default=0)
  for j in range(i+1,n+1):longest[i,j]=max(min(maximum,j-i),longest.get((i+1,j),0))
 @lru_cache(None)
 def paired_score(i,j):
  length=j-i;long=longest[i,j];hits=sum(k>=length for p,k in runs[i]) if length>=4 else 0
  prior={7:2.2,14:2.5,21:1.7,28:1.1,35:.6}.get(length,-.25*min(length%7,7-length%7))
  low_complexity_penalty=6 if q[i:j]=='T'*length else 0
  return .15*length + .4*long - .45*(length-long) + prior - 1 + (1.5 if hits==1 else .5 if hits else 0) - low_complexity_penalty
 def tails(start):
  out=[]
  for length,cost in [(1,.1),(5,1.3),(10,.4)]:
   text=q[:length] if start else q[n-length:]
   if len(text)==length and text=='T'*length:out.append((length,cost))
  if not out:out=[(0,3)]
  return out
 gaps=[(2,.5)] if key=='42hb13' else [(2,.5),(5,.8),(7,2.0)]
 candidates=[]
 for left,lc in tails(True):
  for right,rc in tails(False):
   end=n-right
   @lru_cache(None)
   def best(i):
    possibilities=[]
    if end-i>=3:possibilities.append((paired_score(i,end),(('P',i,end),)))
    for j in range(i+3,end-3):
     for g,cost in gaps:
      if j+g+3>end or q[j:j+g]!='T'*g:continue
      for score,parts in best(j+g):
       possibilities.append((paired_score(i,j)-cost-len(protected.intersection(range(j,j+g)))+score,(('P',i,j),('U',j,j+g))+parts))
    possibilities.sort(key=lambda p:(-p[0],p[1]))
    return tuple(possibilities[:TOP])
   for score,parts in best(left):
    if left:parts=(('U',0,left),)+parts
    if right:parts=parts+(('U',end,n),)
    candidates.append((score-lc-rc,parts))
 candidates.sort(key=lambda p:(-p[0],p[1]));seen=set();unique=[]
 for score,parts in candidates:
  mask=tuple((a,b) for kind,a,b in parts if kind=='U')
  if mask in seen:continue
  seen.add(mask);unique.append((score,parts))
 def detail(candidate):
  score,parts=candidate;items=[]
  for kind,a,b in parts:
   text=q[a:b];hits=[p for p,k in runs[a] if k>=b-a] if kind=='P' and b-a>=4 else []
   items.append({'kind':'paired_candidate' if kind=='P' else 'unpaired_T_candidate','start_1based':a+1,'end_1based':b,'length':b-a,'sequence':text,
                 **({'full_scaffold_hit_count':len(hits),'first_scaffold_indices_0based':hits,'longest_exact_substring_nt':longest[a,b]} if kind=='P' else {})})
  return {'score':round(score,4),'annotated_sequence':''.join(q[a:b] if kind=='P' else '*'+q[a:b]+'*' for kind,a,b in parts),'segments':items}
 winner=detail(unique[0]);alt=[detail(x) for x in unique[1:4]]
 gap=unique[0][0]-unique[1][0] if len(unique)>1 else None
 near=[x for x in unique if unique[0][0]-x[0]<=2.5]
 ambiguous=set()
 for score,parts in near:
  mm={i for kind,a,b in parts if kind=='U' for i in range(a,b)}
  wm={i for kind,a,b in unique[0][1] if kind=='U' for i in range(a,b)}
  ambiguous.update(mm^wm)
 paired=[x for x in winner['segments'] if x['kind']=='paired_candidate']
 total=sum(x['length'] for x in paired)
 support=sum(x['longest_exact_substring_nt'] for x in paired)/total
 notes=[]
 if ambiguous:notes.append('Competing T-boundary assignments within 2.5 score units.')
 if any(x['full_scaffold_hit_count']==0 for x in paired):notes.append('At least one candidate paired block is not a single contiguous p7560 complement; may span a scaffold junction or indicate a wrong boundary.')
 if any(x['length']<6 for x in paired):notes.append('Candidate paired block shorter than the usual REVNANO minimum.')
 if any(x['length']%7 for x in paired):notes.append('Non-7/14-multiple paired lengths retained; possible nick/end offsets.')
 if any(x['kind']=='unpaired_T_candidate' and x['length']==5 and (x['start_1based']==1 or x['end_1based']==n) for x in winner['segments']):notes.append('5T terminal cap is a weaker candidate than the independently supported 10T cap examples.')
 status='ambiguous' if ambiguous else 'clearer estimate' if support>=.85 else 'weak scaffold support'
 assert winner['annotated_sequence'].replace('*','')==q
 assert all(x['sequence']=='T'*x['length'] for x in winner['segments'] if x['kind']=='unpaired_T_candidate')
 return {'original_sequence':q,'status':status,'score_gap':round(gap,4) if gap is not None else None,'ambiguous_positions_1based':[p+1 for p in sorted(ambiguous)],'scaffold_substring_support_fraction':round(support,4),'notes':notes,'proposed':winner,'alternatives':alt}

wb=openpyxl.Workbook();wb.remove(wb.active)
records=[];summary={}
source_workbook=Path('/media/jojo/Archive/NADOC_archive/design_reconstruction/gerling2018/aau1157_Table_S1.xlsx')
original=openpyxl.load_workbook(source_workbook,data_only=True,read_only=True)
headers=['Published ID','Review status','Original sequence','Proposed REVNANO annotation','Colored proposal (red = estimated unpaired T)','Paired block lengths','Unpaired spans (1-based inclusive)','Score margin','Ambiguous base positions','Longest-substring support fraction','Alternative 1','Alternative 2','Notes','Reviewer decision','Reviewer-edited annotation']
for key in ['42hb13','42hb14']:
 ref=json.load(open(R/f'backend/data/gerling2018_{key}.json'))
 expected={i:s for i,s in list(original[ref['sheet']].values)[1:] if isinstance(s,str)}
 assert {x['index']:x['sequence'] for x in ref['sequences']}==expected
 rows=[]
 for source in ref['sequences']:
  row={'sheet':ref['sheet'],'key':key,'published_id':source['index'],**annotate(source['sequence'],key)};rows.append(row)
 records.extend(rows);summary[key]={'count':len(rows),'status_counts':dict(Counter(r['status'] for r in rows)),'proposed_unpaired_T_count':sum(x['length'] for r in rows for x in r['proposed']['segments'] if x['kind']=='unpaired_T_candidate')}
 summary[key]['published_nt_count']=sum(len(r['original_sequence']) for r in rows)
 summary[key]['candidate_paired_nt_count']=summary[key]['published_nt_count']-summary[key]['proposed_unpaired_T_count']
 summary[key]['paired_nt_excess_over_scaffold']=max(0,summary[key]['candidate_paired_nt_count']-N)
 ws=wb.create_sheet(ref['sheet']);ws.append(headers);ws.freeze_panes='F2'
 rev=['# REVIEW DRAFT ONLY — not approved; no solver run performed.','# Star pairs mark ESTIMATED unpaired T bases. Remove stars to recover the exact published oligo.','# Scaffold p7560; numeric comments identify workbook staple IDs.','CIRCULAR',S]
 csvrows=[]
 for row in rows:
  segments=row['proposed']['segments'];colored=CellRichText([TextBlock(InlineFont(rFont='Consolas',color='FFD13B40' if x['kind']=='unpaired_T_candidate' else 'FF234B68',b=x['kind']=='unpaired_T_candidate'),('*'+x['sequence']+'*') if x['kind']=='unpaired_T_candidate' else x['sequence']) for x in segments])
  alts=row['alternatives'];values=[row['published_id'],row['status'],row['original_sequence'],row['proposed']['annotated_sequence'],colored,
    ','.join(str(x['length']) for x in segments if x['kind']=='paired_candidate'),'; '.join(f"{x['start_1based']}-{x['end_1based']} ({x['length']}T)" for x in segments if x['kind']=='unpaired_T_candidate'),row['score_gap'],','.join(map(str,row['ambiguous_positions_1based'])),row['scaffold_substring_support_fraction'],alts[0]['annotated_sequence'] if alts else '',alts[1]['annotated_sequence'] if len(alts)>1 else '', ' '.join(row['notes']),'PENDING','']
  ws.append(values);csvrows.append([str(x) if isinstance(x,CellRichText) else x for x in values]);rev += [f"# Published staple {row['published_id']} | {row['status']}",row['proposed']['annotated_sequence']]
 for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='234B68')
 for rr in ws.iter_rows(min_row=2):
  for cell in rr:cell.alignment=Alignment(vertical='top',wrap_text=True)
  rr[1].fill=PatternFill('solid',fgColor='FFF0C2' if rr[1].value=='ambiguous' else 'E2EDE7')
  ws.row_dimensions[rr[0].row].height=72
 for col in ['C','D','E','K','L','O']:ws.column_dimensions[col].width=70
 for col in ['B','F','G','H','I','J','M','N']:ws.column_dimensions[col].width=26
 ws.auto_filter.ref=ws.dimensions
 with (O/f'{key}_review.csv').open('w',newline='') as f:
  wr=csv.writer(f);wr.writerow(headers);wr.writerows(csvrows)
 (O/f'{key}_estimated_unpaired_T.rev').write_text('\n'.join(rev)+'\n')
 print(key,summary[key],flush=True)
overview=wb.create_sheet('Read me',0)
overview.append(['Review draft — no REVNANO run', 'All original bases retained; *T* blocks are estimates.'])
overview.append(['Method','1T ends, TT crossovers, 7/14-base spacing and p7560 complement evidence; 42hb14 also allows 5T and merged 7T loops.'])
overview.append(['Caution','Scores are rankings, not probabilities. See alternatives and reviewer columns.'])
overview.append(['Global check','Candidate paired totals assume no overlapping occupancy. Excess above 7560 means the draft is not globally feasible.'])
for key,v in summary.items():overview.append([key,f"{v['candidate_paired_nt_count']} candidate paired nt; {v['paired_nt_excess_over_scaffold']} above 7560."])
overview.column_dimensions['A'].width=45
overview.column_dimensions['B'].width=110
for row in overview:
 for cell in row:cell.alignment=Alignment(wrap_text=True,vertical='top')
 overview.row_dimensions[row[0].row].height=45
wb.save(O/'Gerling2018_unpaired_T_review.xlsx')
(O/'annotations.json').write_text(json.dumps(records,indent=2)+'\n')
(O/'summary.json').write_text(json.dumps({'source_workbook_sha256':hashlib.sha256(source_workbook.read_bytes()).hexdigest(),'scaffold_sha256':hashlib.sha256(S.encode()).hexdigest(),'heuristic_version':2,'draft_only':True,'revnano_run':False,'summary':summary},indent=2)+'\n')
# Self-contained reviewer: no network or external JS dependencies.
css='''body{font-family:system-ui;margin:28px;color:#23384b;background:#f6f8fa}header{position:sticky;top:0;background:#f6f8fa;padding:10px 0;border-bottom:1px solid #ccc}article{background:white;padding:16px;margin:14px 0;border:1px solid #ccd5df;border-radius:8px}.seq{font-family:Consolas,monospace;font-size:16px;line-height:2;overflow-wrap:anywhere}.unpaired{color:#b82034;font-weight:bold;background:#ffe4e7}.ambiguous{border-bottom:3px solid #de9b12}.block{display:inline-block;margin-right:4px}.meta{font-size:13px;color:#5b6472}select,input{padding:7px;margin:4px}summary{cursor:pointer}code{overflow-wrap:anywhere}h2{font-size:18px}'''
parts=['<!doctype html><html><head><meta charset="utf-8"><title>Gerling2018 unpaired T review</title><style>'+css+'</style></head><body><h1>Gerling2018 — estimated unpaired T bases</h1><p><b>Review draft. REVNANO has not been rerun.</b> Red starred blocks are proposed unpaired Ts. Orange underlining marks positions whose status changes among near-best alternatives. All original bases are preserved. Paired blocks need not be single contiguous scaffold matches.</p><p>Use the workbook’s Reviewer decision and Reviewer-edited annotation columns to record corrections. Scores are heuristic rankings, not probabilities.</p><header><label>List <select id="sheet"><option value="">Both lists</option><option>42hb (1)-(3)</option><option>42hb (1)-(4)</option></select></label><label>Status <select id="status"><option value="">All</option><option>ambiguous</option><option>clearer estimate</option><option>weak scaffold support</option></select></label><label>Published ID <input id="id" placeholder="e.g. 142" type="number"></label><span id="count"></span></header>']
parts.append('<p><b>Global consistency check:</b> '+ '; '.join(f"{key}: {v['candidate_paired_nt_count']} candidate paired bases ({v['paired_nt_excess_over_scaffold']} above p7560 capacity)" for key,v in summary.items()) + '. Excess means additional unpaired bases or other assumptions must be resolved; these are not solved contact maps.</p>')
for row in records:
 segs=row['proposed']['segments'];amb=set(row['ambiguous_positions_1based']);display=[]
 for x in segs:
  letters=''.join(f'<span class="ambiguous">{base}</span>' if pos in amb else base for pos,base in enumerate(x['sequence'],x['start_1based']))
  un=x['kind']=='unpaired_T_candidate';display.append(f'<span class="block {"unpaired" if un else "paired"}" title="Bases {x["start_1based"]}–{x["end_1based"]}; length {x["length"]}">{"*" if un else ""}{letters}{"*" if un else ""}</span>')
 blocks=','.join(str(x['length']) for x in segs if x['kind']=='paired_candidate')
 parts.append(f'<article data-sheet="{row["sheet"]}" data-status="{row["status"]}" data-id="{row["published_id"]}"><h2>{row["sheet"]} · #{row["published_id"]} · {row["status"]}</h2><div class="seq">'+''.join(display)+f'</div><p class="meta">Paired block lengths: {blocks} · score margin: {row["score_gap"]} · candidate annotations only</p><details><summary>Copyable annotation, alternatives and evidence</summary><p><b>Proposed:</b> <code>{row["proposed"]["annotated_sequence"]}</code></p><p><b>Original:</b> <code>{row["original_sequence"]}</code></p>')
 for i,a in enumerate(row['alternatives'],1):parts.append(f'<p><b>Alternative {i}</b> (score {a["score"]}): <code>{a["annotated_sequence"]}</code></p>')
 parts.append('<p>'+html.escape(' '.join(row['notes']))+'</p><ul>')
 for x in segs:
  if x['kind']=='paired_candidate':parts.append(f'<li>Paired candidate bases {x["start_1based"]}–{x["end_1based"]}: {x["full_scaffold_hit_count"]} full p7560 complement placement(s); longest exact internal match {x["longest_exact_substring_nt"]} nt.</li>')
 parts.append('</ul></details></article>')
parts.append('''<script>const fields=['sheet','status','id'];function filter(){let n=0;document.querySelectorAll('article').forEach(a=>{a.hidden=fields.some(k=>document.getElementById(k).value&&document.getElementById(k).value!==a.dataset[k]);if(!a.hidden)n++});document.getElementById('count').textContent=n+' staples shown'}fields.forEach(k=>document.getElementById(k).addEventListener('input',filter));filter()</script></body></html>''')
(O/'review.html').write_text(''.join(parts))
for r in records:
 if r['key']=='42hb13' and r['published_id'] in [46,142,163,171,174,178]:print(r['published_id'],r['proposed']['annotated_sequence'],r['status'],r['score_gap'])
