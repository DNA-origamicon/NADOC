"""Publish reproducible best-so-far results without changing the NADOC design."""
from pathlib import Path
import ast,csv,json,math,shutil
from collections import Counter
import numpy as np
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.path import Path as MPath
from matplotlib.patches import PathPatch
import openpyxl
from openpyxl.styles import Alignment,Font,PatternFill
R=Path(__file__).resolve().parents[2];O=Path(__file__).resolve().parent;N=7560
best=json.loads((O/'selected_best.json').read_text());results=json.loads((O/'search_results.json').read_text());old=json.loads((R/'analysis/gerling2018_T_annotations_20261027/annotations.json').read_text())
metrics={};maps={};report={};wb=openpyxl.Workbook();overview=wb.active;overview.title='Summary'
overview.append(['Experimental annotation search','Partial maps; not a verified full reconstruction'])
overview.append(['Unique runs',len(results)]);overview.append(['List','Complete staples','Scaffold bases','Uncovered bases','Candidate paired bases','Best trial'])
for key,v in best.items():
 rows=[r for r in old if r['key']==key];a=json.loads((R/v['path']/'annotations.json').read_text());st,sc,circ=json.loads((R/v['path']/'solver/contactmap.json').read_text());m=[];out=[]
 assert circ and len(sc)==N
 for i,row in enumerate(rows):
  placed=str(i) in st
  out.append({'published_id':row['published_id'],'placed':placed,'original_sequence':row['original_sequence'],'approved_annotation':row['proposed']['annotated_sequence'],'best_experimental_annotation':a[i]['annotated_sequence'],'changed':row['proposed']['annotated_sequence']!=a[i]['annotated_sequence'],'paired_block_lengths':','.join(str(s['length']) for s in a[i]['segments'] if s['kind']=='paired_candidate')})
  if placed:
   item=st[str(i)];segments=[];offset=0
   for length in item['section_lengths']:
    indices=item['base_i53'][offset:offset+length]
    if indices[0]>=0:segments.append({'scaffold_indices':indices})
    offset+=length
   m.append({'id':row['published_id'],'segments':segments})
 maps[key]=m;counts=Counter(p for row in m for seg in row['segments'] for p in seg['scaffold_indices'])
 assert len(counts)==v['covered'] and all(c==1 for c in counts.values())
 metrics[key]={'shown':len(m),'total':len(rows),'intervals':sum(len(r['segments']) for r in m),'links':sum(max(0,len(r['segments'])-1) for r in m),'covered':len(counts)}
 report[key]={**v,'total_staples':len(rows),'unplaced_published_ids':[r['published_id'] for r in out if not r['placed']],'uncovered_nt':N-len(counts),'changed_published_ids':[r['published_id'] for r in out if r['changed']]}
 with (O/f'{key}_best_staples.csv').open('w',newline='') as f:
  w=csv.DictWriter(f,fieldnames=list(out[0]));w.writeheader();w.writerows(out)
 with (O/f'{key}_uncovered_regions.csv').open('w',newline='') as f:
  w=csv.writer(f);w.writerow(['start_0based','end_0based_inclusive','length']);start=None
  for p in range(N+1):
   if p<N and p not in counts:
    if start is None:start=p
   elif start is not None:w.writerow([start,p-1,p-start]);start=None
 ws=wb.create_sheet(key);ws.append(list(out[0]))
 for row in out:ws.append(list(row.values()))
 ws.freeze_panes='C2';ws.auto_filter.ref=ws.dimensions
 for col in 'CDE':ws.column_dimensions[col].width=65
 for col in 'ABFG':ws.column_dimensions[col].width=25
 for row in ws:
  for cell in row:cell.alignment=Alignment(wrap_text=True,vertical='top')
  ws.row_dimensions[row[0].row].height=70
 for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='234B68')
 overview.append([key,v['placed'],v['covered'],N-v['covered'],v['candidate_paired_nt'],v['label']])
 shutil.copy2(R/v['path']/'input.rev',O/f'{key}_best.rev')
 shutil.copy2(R/v['path']/'solver/contactmap.csv',O/f'{key}_best_contactmap.csv')
 shutil.copy2(R/v['path']/'solver/contactmap.json',O/f'{key}_best_contactmap.json')
for col in 'ABCDEF':overview.column_dimensions[col].width=35
wb.save(O/'best_annotations_review.xlsx')
(O/'final_report.json').write_text(json.dumps({'runs':len(results),'solver_seconds_total':sum(r['seconds'] for r in results),'best':report},indent=2)+'\n')
with (O/'trial_comparison.csv').open('w',newline='') as f:
 fields=['key','label','placed','covered','candidate_paired_nt','multiple7_blocks','paired_blocks','changed_staples','seconds','solver_error','execution_error','path'];w=csv.DictWriter(f,fieldnames=fields,extrasaction='ignore');w.writeheader();w.writerows(results)
# Reuse the established circle rendering functions without executing its data loading.
source=R/'analysis/gerling2018_circle_graphs_20261027/build_revnano_graphs.py';tree=ast.parse(source.read_text());funcs=[n for n in tree.body if isinstance(n,ast.FunctionDef) and n.name in ['xy','draw','plot']]
exec(compile(ast.Module(body=funcs,type_ignores=[]),str(source),'exec'))
plot([(maps['42hb13'],'42hb13','Best search result · page 3','42hb (1)–(3) · experimental annotations','#b47721'),(maps['42hb14'],'42hb14','Best search result · page 7','42hb (1)–(4) · experimental annotations','#8b579c')],'best_circle_graphs')
native=json.loads((R/'analysis/gerling2018_circle_graphs_20261027/contact_maps.json').read_text())['native']
counts=Counter(p for row in native for seg in row['segments'] for p in seg['scaffold_indices'])
metrics['native']={'shown':len(native),'total':len(native),'intervals':sum(len(r['segments']) for r in native),'links':sum(max(0,len(r['segments'])-1) for r in native),'covered':len(counts)}
plot([(native,'native','Current NADOC design','Stored topology · original scaffold register','#247aa0'),(maps['42hb13'],'42hb13','Best REVNANO search · page 3','Experimental annotations · partial reconstruction','#b47721')],'current_design_vs_best_42hb13')
fig,axs=plt.subplots(1,2,figsize=(12,4.5),layout='constrained')
for ax,key in zip(axs,['42hb13','42hb14']):
 rr=[r for r in results if r['key']==key];maximum=0;progress=[]
 for r in rr:maximum=max(maximum,r['placed']);progress.append(maximum)
 ax.plot(range(1,len(rr)+1),[r['placed'] for r in rr],'.',alpha=.45,label='Trial result');ax.plot(range(1,len(rr)+1),progress,label='Best complete-staple count');ax.axhline(len([r for r in old if r['key']==key]),color='#666',ls='--',label='All published staples');ax.set(title=key,xlabel='Completed trial',ylabel='Complete physical staples');ax.legend(fontsize=8)
fig.savefig(O/'search_progress.png',dpi=150);plt.close(fig)
print(json.dumps({k:{x:v[x] for x in ['placed','covered','candidate_paired_nt','label']} for k,v in best.items()},indent=2))
