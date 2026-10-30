"""Create a schematic, circle graph and linked interactive review of Fig. S1."""
import ast
import csv
import html
import json
import math
from pathlib import Path
import sys

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0,str(ROOT))
OUT = ROOT/'.development-artifacts/gerling2018_trace_reconstruction_20261027'
O = OUT
N = 7560
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.lines import Line2D
from matplotlib.path import Path as MPath
from matplotlib.patches import PathPatch
import numpy as np
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

COLORS = {'unchanged':'#668ba8','corrected':'#9163ad','newly_resolved':'#078087'}


def main():
    global metrics
    records = json.loads((OUT/'matched_staples.json').read_text())
    trace = json.loads((OUT/'figure_trace.json').read_text())
    rows = list(csv.DictReader((OUT/'published_staples.csv').open()))
    by_id = {int(r['published_id']):r for r in rows}
    for r in records:
        r['status'] = by_id[r['published_id']]['status_vs_previous']
        r['color'] = COLORS[r['status']]
    fig,ax = plt.subplots(figsize=(18,12))
    fig.subplots_adjust(left=.055,right=.99,bottom=.1,top=.87)
    for d in trace['scaffold_domains']:
        h = int(d['helix_label'])
        ax.plot([d['start_bp'],d['end_bp']],[h+.22]*2,color='#0073c6',lw=.65,alpha=.8)
    for a,b in zip(trace['scaffold_domains'],trace['scaffold_domains'][1:]):
        ax.plot([a['end_bp'],b['start_bp']],[int(a['helix_label'])+.22,int(b['helix_label'])+.22],color='#0073c6',lw=.45,alpha=.55)
    for r in records:
        for s in r['domains']:
            h = int(s['helix_label'])
            ax.plot([s['lo'],s['hi']],[h-.12]*2,color=r['color'],lw=2.1,solid_capstyle='butt')
        for a,b in zip(r['domains'],r['domains'][1:]):
            x = a['hi'] if a['forward'] else a['lo']
            assert x==(b['lo'] if b['forward'] else b['hi'])
            shift = .65 if a['forward'] else -.65
            ya,yb = int(a['helix_label'])-.12,int(b['helix_label'])-.12
            ax.add_patch(PathPatch(MPath([(x,ya),(x+shift,ya),(x+shift,yb),(x,yb)],
                                        [MPath.MOVETO,MPath.CURVE4,MPath.CURVE4,MPath.CURVE4]),
                                   facecolor='none',edgecolor=r['color'],lw=.65,alpha=.8))
        longest = max(r['domains'],key=lambda s:s['hi']-s['lo'])
        ax.text((longest['lo']+longest['hi'])/2,int(longest['helix_label'])-.30,str(r['published_id']),
                ha='center',va='bottom',fontsize=6,color='#283846')
    ax.scatter([71,72],[41+.22]*2,color='#0073c6',s=10,zorder=4)
    ax.annotate('Scaffold index 0: helix 41, BP 71\n66 unpaired bases follow index 7,493',
                xy=(72,41.22),xytext=(91,42.6),fontsize=9,color='#0066aa',
                arrowprops={'arrowstyle':'->','color':'#0066aa'},ha='left')
    ax.set(xlim=(-12,176),ylim=(43,-1),xlabel='Reconstructed BP coordinate',ylabel='Visible helix label')
    ax.set_yticks(range(42),labels=[str(i) for i in range(42)],fontsize=8)
    ax.set_xticks(range(0,176,7));ax.tick_params(axis='x',labelsize=8)
    ax.grid(axis='x',color='#d9e0e6',lw=.4,alpha=.6)
    ax.legend(handles=[Line2D([0],[0],color=c,lw=3,label=t) for c,t in
                       [(COLORS['newly_resolved'],'27 previously unresolved staples'),
                        (COLORS['corrected'],'44 corrected previous placements'),
                        (COLORS['unchanged'],'142 unchanged placements'),('#0073c6','Traced scaffold')]],
              loc='lower center',bbox_to_anchor=(.5,1.005),ncol=4,fontsize=10)
    fig.suptitle('Gerling Fig. S1 — all 213 published staples reconstructed\n7,494 / 7,494 paired scaffold bases covered; 66-base excess retained',fontsize=18,y=.97)
    fig.text(.5,.025,'Labels on strands are published staple IDs. TT junction inserts and terminal tails are stored in the NADOC design.\nAll full sequences match Table S1 exactly. Coordinates derive from the PDF scaffold and staple paths; no relaxation performed.',ha='center',fontsize=10,color='#536275')
    for ext in ('png','svg'):
        fig.savefig(OUT/f'reconstructed_schematic.{ext}',dpi=170)
    plt.close(fig)

    # Same established renderer as prior circle graphs, with complete true map.
    previous = json.loads((ROOT/'analysis/gerling2018_lattice_completion_20261027/42hb13_best_contactmap.json').read_text())
    actual = json.loads((OUT/'contactmap.json').read_text())
    maps,metrics = {},{}
    for key,contact in [('previous',previous),('complete',actual)]:
        maps[key] = []
        for sid,m in contact[0].items():
            segs,offset = [],0
            for length in m['section_lengths']:
                p = m['base_i53'][offset:offset+length]
                if p[0]>=0:
                    segs.append({'scaffold_indices':p})
                offset += length
            maps[key].append({'id':int(sid)+1,'segments':segs})
        metrics[key] = {'shown':len(maps[key]),'total':213,'covered':sum(b['staple_id']>=0 for b in contact[1].values()),
                        'intervals':sum(len(r['segments']) for r in maps[key]),'links':sum(len(r['segments'])-1 for r in maps[key])}
    source = ROOT/'analysis/gerling2018_circle_graphs_20261027/build_revnano_graphs.py'
    funcs = [n for n in ast.parse(source.read_text()).body if isinstance(n,ast.FunctionDef) and n.name in ('xy','draw')]
    scope = {'np':np,'math':math,'MPath':MPath,'PathPatch':PathPatch,'N':N,'metrics':metrics}
    exec(compile(ast.Module(body=funcs,type_ignores=[]),str(source),'exec'),scope)
    draw = scope['draw']
    fig,axes = plt.subplots(1,2,figsize=(14,8))
    draw(axes[0],maps['previous'],'previous','Previous provisional map','186 staples · connectivity-assisted REVNANO','#668ba8')
    draw(axes[1],maps['complete'],'complete','Complete sequence reconstruction','213 staples · scaffold and staple PDF traces','#078087')
    fig.suptitle('Gerling Fig. S1 — complete staple circle map',fontsize=19,y=.96)
    fig.text(.5,.055,'Rim: p7560 index, clockwise from 0. Interior links connect successive binding intervals of each staple.\nThe only remaining gray interval is indices 7,494–7,559: the 66-base unpaired scaffold segment.',ha='center',fontsize=10)
    fig.text(.5,.015,'All 213 staples and the entire paired region are resolved. The circle is a connectivity view, not the physical shape.',ha='center',fontsize=10,color='#536275')
    fig.subplots_adjust(top=.88,bottom=.13,left=.025,right=.975,wspace=.10)
    for ext in ('png','svg'):
        fig.savefig(OUT/f'complete_circle_graphs.{ext}',dpi=170)
    plt.close(fig)

    # Interactive standalone SVGs share staple selection and a full-sequence table.
    X = lambda bp:55+(bp+10)*5.5
    Y = lambda h:25+int(h)*19
    svg = ['<svg viewBox="0 0 1100 850" role="img" aria-label="Reconstructed schematic">']
    for h in range(42):
        svg += [f'<text x="30" y="{Y(h)+4}" text-anchor="end">{h}</text>',
                f'<line x1="{X(-9)}" x2="{X(173)}" y1="{Y(h)}" y2="{Y(h)}" stroke="#e5e9ed"/>']
    for bp in range(0,176,7):
        svg.append(f'<text x="{X(bp)}" y="830" text-anchor="middle">{bp}</text>')
    for r in records:
        svg.append(f'<g class="staple" data-id="{r["published_id"]}" stroke="{r["color"]}" tabindex="0"><title>Staple {r["published_id"]}; {html.escape(by_id[r["published_id"]]["helix_labels_in_order"])}</title>')
        for s in r['domains']:
            svg.append(f'<line x1="{X(s["lo"])}" x2="{X(s["hi"])}" y1="{Y(s["helix_label"])}" y2="{Y(s["helix_label"])}" stroke-width="5"/>')
        for a,b in zip(r['domains'],r['domains'][1:]):
            x = X(a['hi'] if a['forward'] else a['lo']);y1,y2 = Y(a['helix_label']),Y(b['helix_label']);dx = 4 if a['forward'] else -4
            svg.append(f'<path d="M{x},{y1} C{x+dx},{y1} {x+dx},{y2} {x},{y2}" fill="none" stroke-width="1.1"/>')
        svg.append('</g>')
    svg.append('</svg>')
    circ = ['<svg viewBox="0 0 600 650" role="img" aria-label="Complete circle graph"><circle cx="300" cy="300" r="250" fill="none" stroke="#d9dee5" stroke-width="9"/>']
    def point(p):
        angle = math.pi/2-2*math.pi*p/N
        return 300+250*math.cos(angle),300-250*math.sin(angle)
    for r in records:
        circ.append(f'<g class="staple" data-id="{r["published_id"]}" stroke="{r["color"]}" tabindex="0"><title>Staple {r["published_id"]}</title>')
        segs = maps['complete'][r['published_id']-1]['segments']
        for seg in segs:
            p = seg['scaffold_indices'];pts = ' '.join(f'{x:.2f},{y:.2f}' for x,y in (point(q) for q in p))
            circ.append(f'<polyline points="{pts}" fill="none" stroke-width="6"/>')
        for a,b in zip(segs,segs[1:]):
            x1,y1 = point(a['scaffold_indices'][len(a['scaffold_indices'])//2]);x2,y2 = point(b['scaffold_indices'][len(b['scaffold_indices'])//2])
            circ.append(f'<path d="M{x1},{y1} C{300+(x1-300)*.25},{300+(y1-300)*.25} {300+(x2-300)*.25},{300+(y2-300)*.25} {x2},{y2}" fill="none" stroke-width=".8" opacity=".5"/>')
        circ.append('</g>')
    for p in range(0,N,1000):
        x,y = point(p);x,y = 300+(x-300)*1.11,300+(y-300)*1.11
        circ.append(f'<text x="{x}" y="{y}" text-anchor="middle">{p}</text>')
    circ.append('<text x="300" y="615" text-anchor="middle">7,494 paired + 66 unpaired = 7,560 bases</text></svg>')
    payload = json.dumps({r['published_id']:{**by_id[r['published_id']],'domain_sequences':r['paired_sequences'],
                         'scaffold_routes':r['paired_block_scaffold_indices']} for r in records})
    document = r'''<!doctype html><html lang="en"><meta charset="utf-8"><title>Gerling 2018 reconstruction</title>
<style>body{margin:24px;font:15px system-ui;color:#24384b;background:#fafcfd}h1{font-size:25px;margin-bottom:8px}.panels{display:grid;grid-template-columns:1.7fr 1fr;gap:16px}svg{width:100%;background:white;border:1px solid #ccd7e0;border-radius:8px}svg text{font:11px system-ui;fill:#485e70}label,button,select{font:inherit}button,select{padding:5px}#detail{background:white;padding:16px;border:1px solid #ccd7e0;margin:15px 0;white-space:pre-wrap;overflow-wrap:anywhere;font:14px ui-monospace,monospace;line-height:1.6}.staple{cursor:pointer;transition:opacity .1s}body.selected .staple:not(.chosen){opacity:.08}.chosen{opacity:1;filter:drop-shadow(0 0 1px #333)}.legend span{margin-right:22px}small{color:#536275}@media(max-width:1000px){.panels{grid-template-columns:1fr}}</style>
<h1>Gerling Fig. S1 — 213/213 staples resolved</h1><p>All published sequences match exactly. The full 7,494-base paired region is covered; 66 scaffold bases are unpaired.</p>
<div class="legend"><span style="color:#078087">● 27 newly resolved</span><span style="color:#9163ad">● 44 corrected placements</span><span style="color:#668ba8">● 142 unchanged</span></div>
<p><label>Published staple <select id="pick"><option value="">Show all</option>OPTIONS</select></label> <button id="reset">Clear selection</button> <small>Click either map to highlight the same staple in both views.</small></p>
<div class="panels">SCHEMATIC CIRCLE</div><div id="detail">Select a staple to inspect its helix labels, BP coordinates, full sequence and p7560 binding indices.</div>
<small>Left: corrected physical BP coordinates derived from the PDF, not the old offset-17 projection. Right: scaffold sequence index (0-based). TT junction bases and terminal tails do not occupy paired positions. No relaxation or folding validation performed.</small>
<script>const data=PAYLOAD;const pick=document.querySelector('#pick');function select(id){pick.value=id||'';document.body.classList.toggle('selected',!!id);document.querySelectorAll('.staple').forEach(g=>g.classList.toggle('chosen',g.dataset.id===id));if(!id){document.querySelector('#detail').textContent='Select a staple to inspect its helix labels, BP coordinates, full sequence and p7560 binding indices.';return}const d=data[id];document.querySelector('#detail').textContent=`Staple ${id} · ${d.status_vs_previous.replaceAll('_',' ')}\nHelices: ${d.helix_labels_in_order}\nBP ranges: ${d.domain_bp_ranges}\nPaired domain lengths: ${d.paired_domain_lengths}\nFull sequence: ${d.full_sequence_5to3}\nUnpaired T annotations: ${d.annotated_sequence}\n\nScaffold indices, in staple 5′→3′ order:\n${d.scaffold_routes.map((p,i)=>`Domain ${i+1}: ${p.join(', ')}`).join('\n')}`;}pick.onchange=()=>select(pick.value);document.querySelector('#reset').onclick=()=>select('');document.querySelectorAll('.staple').forEach(g=>{g.onclick=()=>select(g.dataset.id);g.onkeydown=e=>{if(e.key==='Enter')select(g.dataset.id)}});</script></html>'''
    document = document.replace('OPTIONS',''.join(f'<option value="{r["published_id"]}">{r["published_id"]} · {r["status"].replace("_"," ")}</option>' for r in records)).replace('SCHEMATIC',''.join(svg)).replace('CIRCLE',''.join(circ)).replace('PAYLOAD',payload)
    (OUT/'review.html').write_text(document)

    wb = openpyxl.Workbook();ws = wb.active;ws.title='Full published sequences';ws.append(list(rows[0]))
    for row in rows:
        ws.append(list(row.values()))
    ws.freeze_panes='C2';ws.auto_filter.ref=ws.dimensions
    for cell in ws[1]:
        cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='234B68')
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 28
    for col in ('K','L'):
        ws.column_dimensions[col].width=90
    for row in ws.iter_rows(min_row=2):
        ws.row_dimensions[row[0].row].height=70
        for cell in row:
            cell.alignment=Alignment(wrap_text=True,vertical='top')
    wb.save(OUT/'published_staples.xlsx')
    print('Rendered schematic, circles, linked review and workbook.')


if __name__=='__main__':
    main()
